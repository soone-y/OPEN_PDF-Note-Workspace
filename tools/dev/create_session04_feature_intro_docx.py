#!/usr/bin/env python3
"""Create the real Word source used for the Session 04 feature-introduction PDF."""

from __future__ import annotations

import os
from pathlib import Path
from tempfile import NamedTemporaryFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPO_ROOT = Path(__file__).resolve().parents[2]
OUTPUT = REPO_ROOT / "tests/fixtures/office_conversion/Word_PDFノート機能紹介.docx"
ENGLISH_OUTPUT = REPO_ROOT / "tests/fixtures/office_conversion/Word_Feature_Overview.docx"
NAVY = RGBColor(23, 54, 93)
BLUE = RGBColor(46, 116, 181)
INK = RGBColor(31, 41, 55)
MUTED = RGBColor(75, 85, 99)
PALE_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"


def set_run_font(run, *, size: float, color: RGBColor = INK, bold: bool = False) -> None:
    run.font.name = "Yu Gothic"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Yu Gothic")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Yu Gothic")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold


def set_english_title_font(run) -> None:
    """Use a Latin font explicitly so LibreOffice renders the English title."""
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    run.font.size = Pt(24)
    run.font.color.rgb = NAVY
    run.bold = True


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_width(cell, width_dxa: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    width = properties.first_child_found_in("w:tcW")
    if width is None:
        width = OxmlElement("w:tcW")
        properties.append(width)
    width.set(qn("w:w"), str(width_dxa))
    width.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: tuple[int, int]) -> None:
    table.autofit = False
    table_properties = table._tbl.tblPr
    table_width = table_properties.first_child_found_in("w:tblW")
    table_width.set(qn("w:w"), str(sum(widths)))
    table_width.set(qn("w:type"), "dxa")
    table_indent = OxmlElement("w:tblInd")
    table_indent.set(qn("w:w"), "120")
    table_indent.set(qn("w:type"), "dxa")
    table_properties.append(table_indent)
    grid = table._tbl.tblGrid
    for column, width in zip(grid.gridCol_lst, widths):
        column.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_rule(paragraph, color: str = "2E74B5") -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    properties.append(borders)


def add_styled_text(cell, text: str, *, size: float, color: RGBColor = INK, bold: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)


def build_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Yu Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    normal.font.size = Pt(10.5)

    kicker = document.add_paragraph()
    kicker.paragraph_format.space_after = Pt(3)
    run = kicker.add_run("機能紹介 / Office変換サンプル")
    set_run_font(run, size=10, color=BLUE, bold=True)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("PDF Note Workspace")
    set_run_font(run, size=24, color=NAVY, bold=True)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("PDFを読み、ノートと注釈を残し、元データを守るための作業空間")
    set_run_font(run, size=12, color=MUTED)
    add_rule(subtitle)

    lead = document.add_paragraph()
    lead.paragraph_format.space_after = Pt(10)
    run = lead.add_run("このPDFは実際のWordファイルをローカル変換した結果です。")
    set_run_font(run, size=10.5, color=INK, bold=True)

    table = document.add_table(rows=1, cols=2)
    set_table_geometry(table, (2450, 6910))
    header = table.rows[0].cells
    for cell, text in zip(header, ("機能", "できること")):
        set_cell_shading(cell, PALE_BLUE)
        add_styled_text(cell, text, size=10.5, color=NAVY, bold=True)

    rows = (
        ("PDFとノート", "PDFを読みながら、標準ノートやMarkdownなどへ要点を残せます。"),
        ("注釈", "マーカー、文字、図形などをPDFへ重ねて表示します。注釈データはPDFと別に管理します。"),
        ("安全な保存", "原本を直接書き換えず、作業保護とバックアップを使って保存します。"),
        ("Office変換", "通常版ではOfficeファイルをローカルでPDFへ変換し、変換結果を確認できます。"),
    )
    for row_index, (feature, description) in enumerate(rows):
        cells = table.add_row().cells
        if row_index % 2:
            for cell in cells:
                set_cell_shading(cell, LIGHT_GRAY)
        add_styled_text(cells[0], feature, size=10.5, color=NAVY, bold=True)
        add_styled_text(cells[1], description, size=10.2)

    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(14)
    heading.paragraph_format.space_after = Pt(5)
    run = heading.add_run("最初に試すこと")
    set_run_font(run, size=13, color=BLUE, bold=True)

    for text in (
        "第01回のPDFと標準ノートを開き、読みながらメモを残します。",
        "第03回でノート形式とPDFへの往復リンクを試します。",
        "第04回でこのPDFを含むOffice変換結果を確認します。",
    ):
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(text)
        set_run_font(run, size=10.2)

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(8)
    note.paragraph_format.space_after = Pt(0)
    run = note.add_run("注: Lite版にはOffice変換runtimeと第04回の変換PDFは含まれません。")
    set_run_font(run, size=9, color=MUTED)

    properties = document.core_properties
    properties.author = "PDF Note Workspace"
    properties.title = "PDF Note Workspace 機能紹介"
    properties.subject = "Office conversion sample"
    return document


def build_english_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal.font.size = Pt(10.5)

    kicker = document.add_paragraph()
    kicker.paragraph_format.space_after = Pt(3)
    run = kicker.add_run("PDF NOTE WORKSPACE / FEATURE OVERVIEW / OFFICE CONVERSION SAMPLE")
    set_run_font(run, size=10, color=BLUE, bold=True)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("PDF Note Workspace")
    set_english_title_font(run)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("A workspace for reading PDFs, keeping notes and annotations, and protecting original data.")
    set_run_font(run, size=12, color=MUTED)
    add_rule(subtitle)

    lead = document.add_paragraph()
    lead.paragraph_format.space_after = Pt(10)
    run = lead.add_run("This PDF is the result of converting a real Word file locally.")
    set_run_font(run, size=10.5, color=INK, bold=True)

    table = document.add_table(rows=1, cols=2)
    set_table_geometry(table, (2450, 6910))
    header = table.rows[0].cells
    for cell, text in zip(header, ("Feature", "What it supports")):
        set_cell_shading(cell, PALE_BLUE)
        add_styled_text(cell, text, size=10.5, color=NAVY, bold=True)

    rows = (
        ("PDF and notes", "Read a PDF while keeping key points in a standard note or Markdown."),
        ("Annotations", "Show markers, text, and shapes over a PDF. Annotation data is managed separately from the PDF."),
        ("Safe saving", "Keep originals intact while using work protection and backups when saving."),
        ("Office conversion", "In the Full edition, convert Office files to PDF locally and review the result."),
    )
    for row_index, (feature, description) in enumerate(rows):
        cells = table.add_row().cells
        if row_index % 2:
            for cell in cells:
                set_cell_shading(cell, LIGHT_GRAY)
        add_styled_text(cells[0], feature, size=10.5, color=NAVY, bold=True)
        add_styled_text(cells[1], description, size=10.2)

    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(14)
    heading.paragraph_format.space_after = Pt(5)
    run = heading.add_run("Try these first")
    set_run_font(run, size=13, color=BLUE, bold=True)

    for text in (
        "Open the Session 01 PDF and standard note, then keep notes while reading.",
        "Try note formats and the round-trip PDF link in Session 03.",
        "Review this PDF and the Office conversion example in Session 04.",
    ):
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(text)
        set_run_font(run, size=10.2)

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(8)
    note.paragraph_format.space_after = Pt(0)
    run = note.add_run("Note: Lite does not include the Office conversion runtime or the Session 04 conversion PDF.")
    set_run_font(run, size=9, color=MUTED)

    properties = document.core_properties
    properties.author = "PDF Note Workspace"
    properties.title = "PDF Note Workspace Feature Overview"
    properties.subject = "Office conversion sample"
    return document


def main() -> None:
    for output, document in ((OUTPUT, build_document()), (ENGLISH_OUTPUT, build_english_document())):
        output.parent.mkdir(parents=True, exist_ok=True)
        with NamedTemporaryFile(dir=output.parent, suffix=".docx", delete=False) as handle:
            temporary = Path(handle.name)
        try:
            document.save(temporary)
            os.replace(temporary, output)
        finally:
            temporary.unlink(missing_ok=True)
        print(f"created: {output}")


if __name__ == "__main__":
    main()
