#!/usr/bin/env python3
"""Create one-slide PowerPoint fixtures that retain editable native charts."""

from __future__ import annotations

from io import BytesIO
import os
from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile

from lxml import etree
from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[2]
FIXTURE_DIR = REPO_ROOT / "tests/fixtures/office_conversion"
SOURCE = FIXTURE_DIR / "PPTX_グラフ総合検証.pptx"
OUTPUTS = {
    "ja": FIXTURE_DIR / "PPTX_機能紹介とネイティブ図表.pptx",
    "en": FIXTURE_DIR / "PPTX_Feature_Overview_and_Native_Charts.pptx",
}
DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
PRESENTATION_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
RELATIONSHIPS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
CORE_NS = "http://purl.org/dc/elements/1.1/"
NAMESPACES = {"a": DRAWING_NS, "p": PRESENTATION_NS, "c": CHART_NS, "rel": RELATIONSHIPS_NS, "dc": CORE_NS}


def xml_bytes(root: etree._Element) -> bytes:
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def set_slide_text(source: bytes, text: tuple[str, str]) -> bytes:
    root = etree.fromstring(source)
    text_nodes = root.xpath(".//a:t", namespaces=NAMESPACES)
    if len(text_nodes) != 2:
        raise ValueError(f"expected two slide text nodes, found {len(text_nodes)}")
    for node, value in zip(text_nodes, text):
        node.text = value
    return xml_bytes(root)


def expand_native_chart(source: bytes) -> bytes:
    """Use the slide width for the chart while retaining the source PPTX object."""
    root = etree.fromstring(source)
    charts = root.xpath("./p:cSld/p:spTree/p:graphicFrame", namespaces=NAMESPACES)
    if len(charts) != 1:
        raise ValueError(f"expected one native chart frame, found {len(charts)}")
    transform = charts[0].xpath("./p:xfrm", namespaces=NAMESPACES)
    if len(transform) != 1:
        raise ValueError("native chart has no transform")
    offset = transform[0].xpath("./a:off", namespaces=NAMESPACES)
    extent = transform[0].xpath("./a:ext", namespaces=NAMESPACES)
    if len(offset) != 1 or len(extent) != 1:
        raise ValueError("native chart transform is incomplete")
    offset[0].set("x", "457200")
    offset[0].set("y", "777240")
    extent[0].set("cx", "11125200")
    extent[0].set("cy", "4930140")
    return xml_bytes(root)


def set_chart_text(source: bytes, *, title: str, categories: tuple[str, ...], series: tuple[str, ...], language: str) -> bytes:
    root = etree.fromstring(source)
    language_nodes = root.xpath("./c:lang", namespaces=NAMESPACES)
    if len(language_nodes) != 1:
        raise ValueError("native chart has no language declaration")
    language_nodes[0].set("val", language)
    title_nodes = root.xpath("./c:chart/c:title//a:t", namespaces=NAMESPACES)
    if len(title_nodes) != 1:
        raise ValueError(f"expected one native chart title, found {len(title_nodes)}")
    title_nodes[0].text = title
    series_nodes = root.xpath(".//c:barChart/c:ser", namespaces=NAMESPACES)
    if len(series_nodes) != len(series):
        raise ValueError(f"expected {len(series)} chart series, found {len(series_nodes)}")
    for series_node, series_name in zip(series_nodes, series):
        value_node = series_node.xpath("./c:tx/c:strRef/c:strCache/c:pt/c:v", namespaces=NAMESPACES)
        if len(value_node) != 1:
            raise ValueError("native chart series label is missing")
        value_node[0].text = series_name
        category_nodes = series_node.xpath("./c:cat/c:strRef/c:strCache/c:pt/c:v", namespaces=NAMESPACES)
        if len(category_nodes) != len(categories):
            raise ValueError("native chart category cache does not match the sample")
        for category_node, category in zip(category_nodes, categories):
            category_node.text = category
    axis_titles = root.xpath(".//c:catAx/c:title//a:t | .//c:valAx/c:title//a:t", namespaces=NAMESPACES)
    if len(axis_titles) != 2:
        raise ValueError(f"expected two chart axis labels, found {len(axis_titles)}")
    axis_titles[0].text = "カテゴリ" if language == "ja-JP" else "Category"
    axis_titles[1].text = "値" if language == "ja-JP" else "Value"
    return xml_bytes(root)


def set_first_slide_only(source: bytes) -> bytes:
    root = etree.fromstring(source)
    slide_ids = root.xpath("./p:sldIdLst/p:sldId", namespaces=NAMESPACES)
    if not slide_ids:
        raise ValueError("presentation has no slides")
    for slide_id in slide_ids[1:]:
        slide_id.getparent().remove(slide_id)
    return xml_bytes(root)


def remove_unused_slide_relationships(source: bytes) -> bytes:
    root = etree.fromstring(source)
    for relationship in root.xpath("./rel:Relationship", namespaces=NAMESPACES):
        if relationship.get("Type", "").endswith("/slide") and relationship.get("Target") != "slides/slide1.xml":
            relationship.getparent().remove(relationship)
    return xml_bytes(root)


def set_core_title(source: bytes, title: str) -> bytes:
    root = etree.fromstring(source)
    title_nodes = root.xpath("./dc:title", namespaces=NAMESPACES)
    if len(title_nodes) != 1:
        raise ValueError("core properties has no title")
    title_nodes[0].text = title
    return xml_bytes(root)


def set_embedded_chart_data(source: bytes, *, categories: tuple[str, ...], series: tuple[str, ...]) -> bytes:
    workbook = load_workbook(BytesIO(source))
    worksheet = workbook.active
    worksheet.delete_rows(1, worksheet.max_row)
    worksheet.append(["Category", *series])
    values = ((10, 6), (15, 12), (7, 9))
    for category, row_values in zip(categories, values):
        worksheet.append([category, *row_values])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def chart_embedding_path(source: ZipFile) -> str:
    relationships = etree.fromstring(source.read("ppt/charts/_rels/chart1.xml.rels"))
    targets = relationships.xpath("./rel:Relationship[@Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/package']/@Target", namespaces=NAMESPACES)
    if len(targets) != 1:
        raise ValueError("native chart must reference exactly one embedded workbook")
    return "ppt/" + targets[0].replace("../", "")


def create_fixture(locale: str, output: Path) -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError(f"native chart source is missing: {SOURCE}")
    japanese = locale == "ja"
    slide_text = (
        "PDF Note Workspace - Office変換サンプル",
        "PowerPointのネイティブグラフをローカルでPDFへ変換した例。作図機能を示すものではありません。",
    ) if japanese else (
        "Office Conversion Sample",
        "A local PDF conversion of a PowerPoint native chart; it does not indicate that the app creates charts.",
    )
    title = "PowerPointネイティブグラフ（例）" if japanese else "Native PowerPoint Chart (Example)"
    categories = ("日本語 A", "長いカテゴリ名_BBBBB", "C") if japanese else ("Category A", "Long category_BBBBB", "C")
    series = ("系列 1", "系列 2") if japanese else ("Series 1", "Series 2")

    with ZipFile(SOURCE) as source:
        embedding = chart_embedding_path(source)
        replacements = {
            "ppt/presentation.xml": set_first_slide_only(source.read("ppt/presentation.xml")),
            "ppt/_rels/presentation.xml.rels": remove_unused_slide_relationships(source.read("ppt/_rels/presentation.xml.rels")),
            "ppt/slides/slide1.xml": expand_native_chart(
                set_slide_text(source.read("ppt/slides/slide1.xml"), slide_text)
            ),
            "ppt/charts/chart1.xml": set_chart_text(
                source.read("ppt/charts/chart1.xml"), title=title, categories=categories, series=series,
                language="ja-JP" if japanese else "en-US",
            ),
            embedding: set_embedded_chart_data(source.read(embedding), categories=categories, series=series),
            "docProps/core.xml": set_core_title(source.read("docProps/core.xml"), slide_text[0]),
        }
        output.parent.mkdir(parents=True, exist_ok=True)
        with NamedTemporaryFile(dir=output.parent, suffix=".pptx", delete=False) as handle:
            temporary = Path(handle.name)
        try:
            with ZipFile(temporary, "w", ZIP_DEFLATED) as destination:
                for member in source.infolist():
                    destination.writestr(member, replacements.get(member.filename, source.read(member.filename)))
            os.replace(temporary, output)
        finally:
            temporary.unlink(missing_ok=True)
    print(f"created: {output}")


def main() -> None:
    for locale, output in OUTPUTS.items():
        create_fixture(locale, output)


if __name__ == "__main__":
    main()
